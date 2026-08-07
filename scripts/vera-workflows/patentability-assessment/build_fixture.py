from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "references" / "synthetic-patentability-v1"
# The fixture uses the standard_business_brief base font. Long-Chinese Host
# coverage is maintained by the separate Word and litigation acceptance suites.
FONT_NAME = "Calibri"

TARGET_CLAIM = (
    "Claim 1. A winch-assisted drive system for a mountain rescue vehicle, "
    "comprising: a vehicle power source; a first torque output coupled to a "
    "wheel drive path; a second torque output coupled to a winch drum; a "
    "sensor assembly configured to detect cable tension and wheel slip; and "
    "a controller configured, when detected wheel slip exceeds a threshold "
    "and cable tension is below a safety limit, to allocate torque concurrently "
    "to the first torque output and the second torque output."
)

PRIOR_ART_PARAGRAPHS = [
    (
        "PA-P1",
        "A rescue vehicle drive apparatus includes a vehicle power source, a wheel drive path, and a cable drive path coupled to a winch drum.",
    ),
    (
        "PA-P2",
        "A controller switches power between the wheel drive path and the cable drive path according to an operator selection and drives only one path at a time.",
    ),
    (
        "PA-P3",
        "The system monitors cable tension and stops the winch when the tension exceeds a preset safety value.",
    ),
]


def normalize_docx_zip(path):
    temporary = path.with_suffix(".normalized.docx")
    with ZipFile(path, "r") as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    with ZipFile(temporary, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for original, payload in entries:
            normalized = ZipInfo(original.filename, date_time=(2026, 8, 7, 0, 0, 0))
            normalized.compress_type = ZIP_DEFLATED
            normalized.create_system = original.create_system
            normalized.external_attr = original.external_attr
            normalized.flag_bits = original.flag_bits
            target.writestr(normalized, payload)
    temporary.replace(path)


def set_run_font(run, size=11, bold=False, color="000000"):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    properties = run._element.get_or_add_rPr()
    properties.rFonts.set(qn("w:ascii"), FONT_NAME)
    properties.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    properties.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    properties.rFonts.set(qn("w:cs"), FONT_NAME)


def set_style(style, size, color, before, after, bold=True):
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    properties = style._element.get_or_add_rPr()
    properties.rFonts.set(qn("w:ascii"), FONT_NAME)
    properties.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    properties.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    properties.rFonts.set(qn("w:cs"), FONT_NAME)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=9, color="6B7280")


def configure_document(title):
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    set_style(normal, 11, "000000", 0, 6, bold=False)
    normal.paragraph_format.line_spacing = 1.1
    set_style(document.styles["Heading 1"], 16, "2E74B5", 16, 8)
    set_style(document.styles["Heading 2"], 13, "2E74B5", 12, 6)
    set_style(document.styles["Heading 3"], 12, "1F4D78", 8, 4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.add_run("VERA SYNTHETIC PATENTABILITY QA"), size=9, bold=True, color="6B7280")
    footer = section.footer.paragraphs[0]
    set_run_font(footer.add_run("Synthetic QA fixture • No real client • Not for filing  |  "), size=9, color="6B7280")
    add_page_number(footer)

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(16)
    title_paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(title_paragraph.add_run(title), size=23, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run_font(
        subtitle.add_run("PAT-DEMO-PAT-01 • internal workflow acceptance source"),
        size=13,
        color="374151",
    )
    return document


def add_label(document, label, value):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(f"{label}: "), bold=True)
    set_run_font(paragraph.add_run(value))


def build_target():
    document = configure_document("Synthetic Target Claim and Invention Disclosure")
    add_label(document, "Fixture ID", "PAT-DEMO-PAT-01-TARGET")
    add_label(document, "Matter role", "Target claim / invention disclosure")
    add_label(document, "Classification", "Entirely synthetic QA material; no real applicant, inventor, product, or filing")
    add_label(document, "Workflow boundary", "Source only; the original must remain immutable")

    document.add_heading("Use limitation", level=1)
    document.add_paragraph(
        "This document exists only to test Matter ownership, fixed DocumentVersion inputs, feature decomposition, citations, unresolved gaps, and editable work-product generation. It is not a patent application, legal authority, prior-art search, or legal opinion."
    )
    document.add_heading("Target claim", level=1)
    claim = document.add_paragraph()
    claim.paragraph_format.keep_together = True
    set_run_font(claim.add_run("TC-C1 — "), bold=True, color="1F4D78")
    set_run_font(claim.add_run(TARGET_CLAIM))

    document.add_heading("Invention disclosure context", level=1)
    context = document.add_paragraph()
    set_run_font(context.add_run("TD-P1 — "), bold=True, color="1F4D78")
    set_run_font(
        context.add_run(
            "The test system is intended for mountain rescue on low-traction surfaces. Its design objective is to coordinate wheel traction and winch traction when the vehicle has material wheel slip while cable tension remains within a safe range."
        )
    )
    gap = document.add_paragraph()
    set_run_font(gap.add_run("TD-P2 — "), bold=True, color="1F4D78")
    set_run_font(
        gap.add_run(
            "This synthetic disclosure does not provide threshold values, a control period, sensor calibration, or alternative embodiments. Those points must remain unresolved and must not be supplied by a model."
        )
    )
    return document


def build_prior_art():
    document = configure_document("Synthetic Prior-Art Publication Excerpt")
    add_label(document, "Fixture ID", "PAT-DEMO-PAT-01-PA1")
    add_label(document, "Matter role", "Pinned prior-art source")
    add_label(document, "Synthetic identifier", "VERA-SYNTHETIC-PA-001 (not an official publication number)")
    add_label(document, "Classification", "Entirely synthetic QA material; no real publication or assignee")
    add_label(document, "Review scope", "The three numbered paragraphs below are the complete fixture text")

    document.add_heading("Source notice", level=1)
    document.add_paragraph(
        "This source is intentionally synthetic and cannot prove real-world patentability, validity, infringement, or search completeness. A workflow output must preserve that limitation and require verification against an official publication before any external reliance."
    )
    document.add_heading("Disclosure excerpt", level=1)
    for identifier, text in PRIOR_ART_PARAGRAPHS:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        set_run_font(paragraph.add_run(f"{identifier} — "), bold=True, color="1F4D78")
        set_run_font(paragraph.add_run(text))

    document.add_heading("Explicit coverage gap", level=1)
    gap = document.add_paragraph()
    set_run_font(gap.add_run("PA-G1 — "), bold=True, color="9B1C1C")
    set_run_font(
        gap.add_run(
            "The fixture has no official publication PDF, publication date, priority date, family record, or additional embodiments. Treat source verification and search coverage as unresolved."
        )
    )
    return document


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    outputs = [
        (OUTPUT / "PAT-DEMO-PAT-01-target-claim.docx", build_target()),
        (OUTPUT / "PAT-DEMO-PAT-01-synthetic-prior-art.docx", build_prior_art()),
    ]
    for path, document in outputs:
        document.core_properties.title = path.stem
        document.core_properties.subject = "Synthetic Vera patentability workflow QA fixture"
        document.core_properties.author = "Vera QA"
        document.core_properties.keywords = "synthetic, QA, patentability, not for filing"
        document.core_properties.comments = "No real client, inventor, assignee, publication, or legal conclusion."
        document.core_properties.created = datetime(2026, 8, 7, 0, 0, 0)
        document.core_properties.modified = datetime(2026, 8, 7, 0, 0, 0)
        document.save(path)
        normalize_docx_zip(path)
        print(path)


if __name__ == "__main__":
    main()
